"""File ingest: CSV/XLSX -> canonical DataFrames plus issues (TR-07..TR-12, PRD 4.1)."""

from __future__ import annotations

import csv
import io
import math
import re
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from statistics import median
from typing import Any, Literal

import pandas as pd

from engine.schema import LOG_COLS, Issue

ALIASES: dict[str, str] = {
    "terminal_id": "tid",
    "atm_id": "tid",
    "id_atm": "tid",
    "tanggal": "date",
    "tgl": "date",
    "log_date": "date",
    "saldo_awal": "opening_balance",
    "open_balance": "opening_balance",
    "penarikan": "withdrawal",
    "tarik_tunai": "withdrawal",
    "cash_out": "withdrawal",
    "isi_ulang": "replenishment",
    "pengisian": "replenishment",
    "cash_in": "replenishment",
    "refill": "replenishment",
    "saldo_akhir": "closing_balance",
    "close_balance": "closing_balance",
    "end_balance": "closing_balance",
}

MASTER_COLS = [
    "tid",
    "location_name",
    "location_type",
    "capacity",
    "min_safe_balance",
    "replenish_cycle_days",
    "denomination",
    "cost_per_visit",
    "current_fill_level",
    "next_visit_date",
]
MASTER_INT = [
    "capacity",
    "min_safe_balance",
    "replenish_cycle_days",
    "denomination",
    "cost_per_visit",
    "current_fill_level",
]
MASTER_REQUIRED = ["tid", *MASTER_INT]
HOLIDAY_COLS = ["date", "name", "type"]
LOG_REQUIRED = ["tid", "date", "opening_balance", "closing_balance"]
LOG_AMOUNTS = ["opening_balance", "withdrawal", "replenishment", "closing_balance"]
OPTIONAL_AMOUNTS = ("withdrawal", "replenishment")  # PRD 4.1: derivable from the balances when missing

_AMOUNT_RE = re.compile(r"^\s*(Rp)?\s*-?\d+(\.\d+)?\s*$", re.IGNORECASE)
# id-ID thousand grouping (one or more groups of 3), e.g. 1.000 or 250.000.000: ambiguous, never guessed.
_THOUSAND_RE = re.compile(r"^\s*(Rp)?\s*-?\d{1,3}(\.\d{3})+\s*$", re.IGNORECASE)
_ISO_RE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")
_DMY_RE = re.compile(r"^(\d{2})[/-](\d{2})[/-](\d{4})$")


class UnreadableFile(Exception):
    def __init__(self, issues: list[Issue]):
        self.issues = issues
        super().__init__("; ".join(i.message for i in issues))


def _normalise_header(h: Any) -> str:
    return re.sub(r"[\s-]+", "_", str(h if h is not None else "").strip().lower())


def _read_rows(data: bytes, filename: str) -> list[list[Any]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise UnreadableFile([Issue("error", "UNREADABLE", f"Cannot open XLSX: {e}")]) from e
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise UnreadableFile([Issue("error", "BAD_ENCODING", "File is not valid UTF-8")]) from e
    sample = text[:4096]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;").delimiter
    except csv.Error:
        lines = sample.splitlines()
        first = lines[0] if lines else ""
        delim = ";" if first.count(";") > first.count(",") else ","
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]


def _parse_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not isinstance(v, str):
        return None
    s = v.strip()
    try:
        if m := _ISO_RE.match(s):
            return date(int(m[1]), int(m[2]), int(m[3]))
        if m := _DMY_RE.match(s):
            return date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None
    return None


def _parse_amount(v: Any) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None
        return int(Decimal(str(v)).quantize(Decimal(1), rounding=ROUND_HALF_UP))
    # if not isinstance(v, str) or not _AMOUNT_RE.match(v):
    if not isinstance(v, str) or not _AMOUNT_RE.match(v) or _THOUSAND_RE.match(v):
        return None
    num = re.sub(r"(?i)rp|\s", "", v)
    return int(Decimal(num).quantize(Decimal(1), rounding=ROUND_HALF_UP))


# def _parse_amount_cols(r, cols, cell, issues: list[Issue], tid: str, d: date | None = None):
def _parse_amount_cols(
    r, cols, cell, issues: list[Issue], tid: str, d: date | None = None, optional: tuple[str, ...] = ()
):
    # """Parse the named amount columns of one row; return {col: int} or None after recording the issue."""
    """Parse the named amount columns of one row; return {col: int} or None after recording the issue.
    A blank cell of an `optional` column is None (missing for that row: derived later, F6)."""
    # rec: dict[str, int] = {}
    rec: dict[str, int | None] = {}
    for c in cols:
        raw = cell(r, c)
        if c in optional and _is_blank(raw):
            rec[c] = None
            continue
        v = None if _is_blank(raw) else _parse_amount(raw)
        if v is None:
            issues.append(Issue("error", "BAD_AMOUNT", f"Unparseable amount {raw!r} in {c}", tid, d))
            return None
        if v < 0:
            issues.append(Issue("error", "NEGATIVE_VALUE", f"Negative value in {c}", tid, d))
            return None
        rec[c] = v
    return rec


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def read_table(
    data: bytes,
    filename: str,
    kind: Literal["log", "master", "holidays"],
    detect_jump_frac: float = 0.20,
    capacities: dict[str, int] | None = None,
) -> tuple[pd.DataFrame, list[Issue]]:
    rows = _read_rows(data, filename)
    rows = [r for r in rows if any(not _is_blank(c) for c in r)]
    if not rows:
        raise UnreadableFile([Issue("error", "EMPTY_FILE", "File has no rows")])
    raw_header = [_normalise_header(h) for h in rows[0]]
    if kind == "log":
        header = [ALIASES.get(h, h) for h in raw_header]
        known, required = LOG_COLS, LOG_REQUIRED
    elif kind == "master":
        header, known, required = raw_header, MASTER_COLS, MASTER_REQUIRED
    else:
        header, known, required = raw_header, HOLIDAY_COLS, HOLIDAY_COLS

    issues: list[Issue] = []
    missing = [c for c in required if c not in header]
    if missing:
        raise UnreadableFile(
            [Issue("error", "MISSING_COLUMN", f"Required column '{c}' is missing") for c in missing]
        )
    for h, orig in zip(header, rows[0], strict=False):
        if h not in known:
            issues.append(Issue("info", "UNKNOWN_COLUMN", f"Column '{orig}' ignored"))

    # First occurrence of each canonical column wins.
    idx: dict[str, int] = {}
    for i, h in enumerate(header):
        if h in known and h not in idx:
            idx[h] = i

    def cell(r: list[Any], col: str) -> Any:
        i = idx.get(col)
        return r[i] if i is not None and i < len(r) else None

    if kind == "log":
        df = _read_log(rows[1:], idx, cell, issues)
        df = _aggregate(df, issues)
        return _finish_log(df, issues, detect_jump_frac, capacities or {})
    if kind == "master":
        return _read_master(rows[1:], cell, issues), issues
    return _read_holidays(rows[1:], cell, issues), issues


def _read_log(body, idx, cell, issues: list[Issue]) -> pd.DataFrame:
    out: list[dict[str, Any]] = []
    amount_cols = [c for c in LOG_AMOUNTS if c in idx]
    for r in body:
        tid = cell(r, "tid")
        tid = "" if _is_blank(tid) else str(tid).strip()
        d = _parse_date(cell(r, "date"))
        if d is None:
            issues.append(Issue("error", "BAD_DATE", f"Unparseable date {cell(r, 'date')!r}", tid or None))
            continue
        if not tid:
            issues.append(Issue("error", "BAD_TID", "Missing tid", None, d))
            continue
        # rec: dict[str, Any] = {"tid": tid, "date": d}
        # ok = True
        # for c in amount_cols:
        #     raw = cell(r, c)
        #     v = None if _is_blank(raw) else _parse_amount(raw)
        #     if v is None:
        #         issues.append(Issue("error", "BAD_AMOUNT", f"Unparseable amount {raw!r} in {c}", tid, d))
        #         ok = False
        #         break
        #     if v < 0:
        #         issues.append(Issue("error", "NEGATIVE_VALUE", f"Negative value in {c}", tid, d))
        #         ok = False
        #         break
        #     rec[c] = v
        # if ok:
        #     out.append(rec)
        # amounts = _parse_amount_cols(r, amount_cols, cell, issues, tid, d)
        amounts = _parse_amount_cols(r, amount_cols, cell, issues, tid, d, OPTIONAL_AMOUNTS)
        if amounts is not None:
            out.append({"tid": tid, "date": d, **amounts})
    return pd.DataFrame(out, columns=["tid", "date", *amount_cols])


def _aggregate(df: pd.DataFrame, issues: list[Issue]) -> pd.DataFrame:
    if df.empty:
        return df
    dup = df.duplicated(["tid", "date"], keep=False)
    if not dup.any():
        return df
    for (tid, d), _ in df[dup].groupby(["tid", "date"], sort=True):
        issues.append(Issue("info", "DUPLICATE_KEY", "Multiple rows aggregated", tid, d))
    spec: dict[str, str] = {"opening_balance": "first", "closing_balance": "last"}
    for c in ("withdrawal", "replenishment"):
        if c in df.columns:
            # spec[c] = "sum"
            spec[c] = _sum_or_missing
    return df.groupby(["tid", "date"], sort=True, as_index=False).agg(spec)


def _sum_or_missing(s: pd.Series) -> float:
    """Aggregated amount; missing (NaN) when every duplicate row left the cell blank (F6)."""
    return s.sum(min_count=1)


def _finish_log(df, issues, frac, capacities):
    # derive = "withdrawal" not in df.columns or "replenishment" not in df.columns
    has_w, has_r = "withdrawal" in df.columns, "replenishment" in df.columns
    derive = not (has_w and has_r)
    # if derive:
    if not has_w and not has_r:
        df = _derive(df, frac, capacities)
    elif not has_r:
        df = df.copy()
        df["replenishment"] = (df["closing_balance"] - df["opening_balance"] + df["withdrawal"]).clip(lower=0)
    elif not has_w:
        df = df.copy()
        df["withdrawal"] = df["opening_balance"] + df["replenishment"] - df["closing_balance"]
    df = df.copy()
    df["source"] = "derived" if derive else "upload"
    df = _derive_blank_cells(df, issues, frac, capacities)
    df = df[[*LOG_COLS, "source"]].sort_values(["tid", "date"]).reset_index(drop=True)
    for c in LOG_COLS[2:]:
        df[c] = df[c].astype("int64")
    return df, issues


def _derive_blank_cells(df: pd.DataFrame, issues: list[Issue], frac: float, capacities: dict[str, int]):
    """F6: per-row PRD 4.1 derivation of blank withdrawal/replenishment cells (columns present)."""
    miss_w, miss_r = df["withdrawal"].isna(), df["replenishment"].isna()
    if not (miss_w | miss_r).any():
        return df
    df = df.copy()
    for c in ("withdrawal", "replenishment"):
        df[c] = pd.to_numeric(df[c]).astype("float64")
    op, cl = df["opening_balance"], df["closing_balance"]
    only_r = miss_r & ~miss_w
    df.loc[only_r, "replenishment"] = (cl - op + df["withdrawal"]).clip(lower=0)[only_r]
    only_w = miss_w & ~miss_r
    df.loc[only_w, "withdrawal"] = (op + df["replenishment"] - cl)[only_w]
    both = miss_w & miss_r
    if both.any():  # neither flow known: balance-jump detection over the TID's balances
        est = _derive(df[["tid", "date", "opening_balance", "closing_balance"]], frac, capacities)
        est = est.set_index(["tid", "date"]).reindex(pd.MultiIndex.from_frame(df[["tid", "date"]]))
        for c in ("withdrawal", "replenishment"):
            df.loc[both, c] = est[c].to_numpy()[both.to_numpy()]
    derived = miss_w | miss_r
    df.loc[derived, "source"] = "derived"
    for (tid, col), g in pd.concat(
        [
            df.loc[miss_w, ["tid", "date"]].assign(col="withdrawal"),
            df.loc[miss_r, ["tid", "date"]].assign(col="replenishment"),
        ]
    ).groupby(["tid", "col"], sort=True):
        n = len(g)
        issues.append(
            Issue(
                "info",
                "DERIVED_VALUE",
                f"{col} blank on {n} row{'s' if n != 1 else ''}: derived from the balances (PRD 4.1)",
                str(tid),
                min(g["date"]),
            )
        )
    for tid, d, w in zip(df["tid"][derived], df["date"][derived], df["withdrawal"][derived], strict=True):
        if w < 0:
            issues.append(
                Issue("warning", "NEGATIVE_DERIVED", f"Derived withdrawal is negative ({int(w)})", tid, d)
            )
    return df


def _derive(df: pd.DataFrame, frac: float, capacities: dict[str, int]) -> pd.DataFrame:
    """PRD 4.1: a balance jump > frac * capacity marks a replenishment day."""
    out_rows: list[dict[str, Any]] = []
    for tid, g in df.groupby("tid", sort=True):
        g = g.sort_values("date")
        cap = capacities.get(tid)
        if cap is None:
            cap = int(max(g["opening_balance"].max(), g["closing_balance"].max()))
        wd_by_date: dict[date, int] = {}
        for rec in g[["tid", "date", "opening_balance", "closing_balance"]].to_dict("records"):
            d = rec["date"]
            diff = rec["closing_balance"] - rec["opening_balance"]
            if diff > frac * cap:
                same_wd = [
                    wd_by_date[d - timedelta(days=7 * k)]
                    for k in range(1, 5)
                    if (d - timedelta(days=7 * k)) in wd_by_date
                ]
                if same_wd:
                    est = median(same_wd)
                elif wd_by_date:
                    est = median(wd_by_date.values())
                else:
                    est = 0
                est = round(est)
                wd, rep = est, diff + est
            else:
                wd, rep = rec["opening_balance"] - rec["closing_balance"], 0
            wd_by_date[d] = wd
            out_rows.append({**rec, "withdrawal": wd, "replenishment": rep})
    return pd.DataFrame(out_rows, columns=LOG_COLS)


def _read_master(body, cell, issues: list[Issue]) -> pd.DataFrame:
    out: dict[str, dict[str, Any]] = {}
    for r in body:
        tid = cell(r, "tid")
        tid = "" if _is_blank(tid) else str(tid).strip()
        if not tid:
            issues.append(Issue("error", "BAD_TID", "Missing tid"))
            continue
        # rec: dict[str, Any] = {"tid": tid}
        # ok = True
        # for c in MASTER_INT:
        #     raw = cell(r, c)
        #     v = None if _is_blank(raw) else _parse_amount(raw)
        #     if v is None:
        #         issues.append(Issue("error", "BAD_AMOUNT", f"Unparseable value {raw!r} in {c}", tid))
        #         ok = False
        #         break
        #     if v < 0:
        #         issues.append(Issue("error", "NEGATIVE_VALUE", f"Negative value in {c}", tid))
        #         ok = False
        #         break
        #     rec[c] = v
        # if not ok:
        #     continue
        amounts = _parse_amount_cols(r, MASTER_INT, cell, issues, tid)
        if amounts is None:
            continue
        rec: dict[str, Any] = {"tid": tid, **amounts}
        for c, default in (("location_name", tid), ("location_type", "unknown")):
            raw = cell(r, c)
            rec[c] = default if _is_blank(raw) else str(raw).strip()
        raw = cell(r, "next_visit_date")
        nv = None
        if not _is_blank(raw):
            nv = _parse_date(raw)
            if nv is None:
                issues.append(Issue("error", "BAD_DATE", f"Unparseable date {raw!r}", tid))
                continue
        rec["next_visit_date"] = nv
        if tid in out:
            issues.append(Issue("info", "DUPLICATE_KEY", "Duplicate tid; last row kept", tid))
        out[tid] = rec
    df = pd.DataFrame(list(out.values()), columns=MASTER_COLS)
    for c in MASTER_INT:
        df[c] = df[c].astype("int64")
    df["next_visit_date"] = df["next_visit_date"].astype(object)
    return df


def _read_holidays(body, cell, issues: list[Issue]) -> pd.DataFrame:
    out = []
    for r in body:
        d = _parse_date(cell(r, "date"))
        if d is None:
            issues.append(Issue("error", "BAD_DATE", f"Unparseable date {cell(r, 'date')!r}"))
            continue
        name, typ = cell(r, "name"), cell(r, "type")
        out.append(
            {
                "date": d,
                "name": "" if _is_blank(name) else str(name).strip(),
                "type": "" if _is_blank(typ) else str(typ).strip(),
            }
        )
    return pd.DataFrame(out, columns=HOLIDAY_COLS)
